# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Redactor contributors
from pathlib import Path
import pytest
from redactor.formats import export_file, read_file, tesseract_path


@pytest.mark.parametrize("extension", [".txt", ".md", ".csv", ".tsv", ".pdf", ".docx", ".xlsx", ".pptx", ".rtf", ".odt", ".ods", ".odp"])
def test_content_export_import(extension, tmp_path, qapp):
    path = tmp_path / ("work" + extension)
    export_file(path, "Leia Organa R1234ABCD\n203.0.113.8\nReview completed.")
    assert path.stat().st_size > 0
    result = read_file(path).text
    assert "Leia Organa R1234ABCD" in result
    assert "203.0.113.8" in result


@pytest.mark.parametrize("extension", [".png", ".jpg", ".tiff", ".bmp", ".webp"])
def test_image_export(extension, tmp_path):
    from PIL import Image
    path = tmp_path / ("work" + extension)
    export_file(path, "A reviewed work product")
    with Image.open(path) as image:
        assert image.width == 1400
        # TIFF uses this structure for mandatory image dimensions/compression too.
        assert 315 not in image.getexif()  # Artist
        assert 34853 not in image.getexif()  # GPS metadata


def test_local_ocr(tmp_path):
    if not tesseract_path(): pytest.skip("Local Tesseract unavailable")
    path = tmp_path / "scan.png"
    export_file(path, "Sensitive account 123-45-6789")
    assert "123-45-6789" in read_file(path).text


def test_spreadsheet_formula_is_literal(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "work.xlsx"
    export_file(path, '=HYPERLINK("https://example.com")')
    book = load_workbook(path)
    assert book.active["A1"].data_type == "s"
    book.close()


def test_csv_formula_guard(tmp_path):
    import csv
    path = tmp_path / "work.csv"
    export_file(path, "=1+1\t@SUM(1)\tordinary")
    with path.open(encoding="utf-8-sig") as f:
        assert next(csv.reader(f)) == ["'=1+1", "'@SUM(1)", "ordinary"]


def test_source_metadata_does_not_survive_export(tmp_path):
    from docx import Document
    import zipfile
    original, output = tmp_path / "original.docx", tmp_path / "output.docx"
    doc = Document()
    doc.core_properties.author = "SENSITIVE AUTHOR"
    doc.add_paragraph("Jane Smith")
    doc.save(original)
    assert "Jane Smith" in read_file(original).text
    export_file(output, "Leia Organa R1234ABCD")
    with zipfile.ZipFile(output) as archive:
        raw = b"".join(archive.read(name) for name in archive.namelist())
    assert b"SENSITIVE AUTHOR" not in raw and b"Jane Smith" not in raw


def test_no_silent_image_truncation(tmp_path):
    with pytest.raises(ValueError, match="multiple image pages"):
        export_file(tmp_path / "long.png", "Line\n" * 100)


def test_legacy_format_explicit_error(tmp_path):
    with pytest.raises(ValueError, match="legacy"):
        read_file(tmp_path / "old.doc")


def test_xlsx_long_cell_not_truncated(tmp_path):
    with pytest.raises(ValueError, match="Excel"):
        export_file(tmp_path / "long.xlsx", "a" * 40000)
def test_pdf_export_creates_nested_destination(tmp_path, qapp):
    path = tmp_path / 'new' / 'exports' / 'work.pdf'
    export_file(path, 'Nested PDF destination contains this text.')
    assert 'Nested PDF destination' in read_file(path).text


def test_pdf_roundtrip_does_not_add_page_numbers(tmp_path, qapp):
    text = "\n".join(f"Evidence row {number:03d}: IBM XYZ" for number in range(150))
    path = tmp_path / 'paginated.pdf'
    export_file(path, text)
    assert read_file(path).text.splitlines() == text.splitlines()
