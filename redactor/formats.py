# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Redactor contributors
"""Read content locally; export fresh text-only artifacts, never original containers."""
from __future__ import annotations

import csv
import io
import os
import shutil
import subprocess
import sys
import textwrap
import zipfile
from dataclasses import dataclass
from pathlib import Path

MAX_BYTES = 50 * 1024 * 1024
MAX_TEXT = 1_000_000
INPUT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".log", ".html", ".pdf", ".docx", ".xlsx", ".pptx", ".rtf", ".odt", ".ods", ".odp", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
OUTPUT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".pdf", ".docx", ".xlsx", ".pptx", ".rtf", ".odt", ".ods", ".odp", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
NOTICE = "Content-only import/export: original layout, images, attachments, formulas, comments and metadata are not preserved. Review extracted text, including OCR errors and omitted visual context."


@dataclass
class Imported:
    text: str
    note: str = NOTICE


def tesseract_path() -> str | None:
    from .portable import root
    candidates = [root() / 'tools/tesseract/tesseract.exe', root() / 'tools/tesseract/bin/tesseract', root() / 'tools/tesseract/tesseract']
    return next((str(p) for p in candidates if p.is_file()), None)


def ocr_image(image) -> str:
    executable = tesseract_path()
    if not executable:
        raise ValueError("Image/scanned-page import needs local Tesseract OCR. See Help → User Guide for your platform's setup instructions. No data was uploaded.")
    buffer = io.BytesIO()
    image.convert("RGB").save(buffer, format="PNG")
    env = os.environ.copy()
    from .portable import root
    env['TESSDATA_PREFIX'] = str(root() / 'tools/tesseract/tessdata')
    env['LD_LIBRARY_PATH'] = str(root() / 'tools/tesseract/lib')
    result = subprocess.run([executable, "stdin", "stdout", "-l", "eng"], input=buffer.getvalue(),
                            capture_output=True, timeout=120, env=env,
                            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    if result.returncode:
        if "--smoke-test" in sys.argv:
            print("Packaged OCR diagnostic:", result.stderr.decode("utf-8", errors="replace"), flush=True)
        raise ValueError(f"R004: Local OCR exited with code {result.returncode}. Check the bundled tools/tesseract executable, libraries and English language data. See Help → Error Guide, R004.")
    return result.stdout.decode("utf-8", errors="replace").strip()


def check_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        items = archive.infolist()
        if len(items) > 20000 or sum(i.file_size for i in items) > 150 * 1024 * 1024:
            raise ValueError("This document expands beyond the 150 MB safety limit.")


def read_file(path: Path) -> Imported:
    from PIL import Image, ImageOps, ImageSequence
    extension = path.suffix.lower()
    if extension not in INPUT_EXTENSIONS:
        raise ValueError("Unsupported format. Convert legacy .doc/.xls/.ppt files to .docx/.xlsx/.pptx in Office first.")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Import limit is 50 MB per file. Split the document first.")
    if extension in {".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp"}:
        check_archive(path)
    parts: list[str] = []
    if extension == ".pdf":
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(path)) as doc:
            if len(doc) > 500:
                raise ValueError("PDF limit is 500 pages. Split the document first.")
            for index in range(len(doc)):
                page = doc[index]
                try:
                    text_page = page.get_textpage()
                    try:
                        extracted = text_page.get_text_range().replace("\r\n", "\n").strip()
                    finally:
                        text_page.close()
                    if not extracted:
                        width, height = page.get_size()
                        if width * height * 4 > 40_000_000:
                            raise ValueError("PDF page is too large to rasterize safely. Resize it before importing.")
                        bitmap = page.render(scale=2)
                        try:
                            extracted = ocr_image(bitmap.to_pil())
                        finally:
                            bitmap.close()
                    parts.append(extracted)
                finally:
                    page.close()
    elif extension == ".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc = Document(path)
        def blocks(container):
            for block in container.iter_inner_content():
                if isinstance(block, Paragraph):
                    parts.append(block.text)
                elif isinstance(block, Table):
                    for row in block.rows:
                        parts.append("\t".join(cell.text for cell in row.cells))
        for section in doc.sections:
            for area in (section.header, section.first_page_header, section.even_page_header):
                blocks(area)
        blocks(doc)
        for section in doc.sections:
            for area in (section.footer, section.first_page_footer, section.even_page_footer):
                blocks(area)
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book:
                parts.append(sheet.title)
                for row in sheet.iter_rows():
                    parts.append("\t".join(str(c.value) if c.value is not None else "" for c in row))
                    if sum(len(p) for p in parts[-100:]) > MAX_TEXT or len(parts) > 100000:
                        raise ValueError("Worksheet exceeds the import limit.")
        finally:
            book.close()
    elif extension == ".pptx":
        from pptx import Presentation
        def shapes(items):
            for shape in items:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend("\t".join(c.text for c in row.cells) for row in shape.table.rows)
                if hasattr(shape, "shapes"):
                    shapes(shape.shapes)
        for slide in Presentation(path).slides:
            shapes(slide.shapes)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
    elif extension in {".odt", ".ods", ".odp"}:
        from odf.opendocument import load
        from odf import teletype
        from odf.text import P, H
        doc = load(str(path))
        parts = [teletype.extractText(element) for element in doc.getElementsByType(P)]
        parts += [teletype.extractText(element) for element in doc.getElementsByType(H)]
    elif extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        with Image.open(path) as source:
            for i, frame in enumerate(ImageSequence.Iterator(source)):
                if i >= 100:
                    raise ValueError("Image limit is 100 frames.")
                parts.append(ocr_image(ImageOps.exif_transpose(frame)))
    else:
        data = path.read_bytes()
        try:
            text = data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("Save the text file as UTF-8 or UTF-16 before importing.") from error
        if extension == ".rtf":
            from striprtf.striprtf import rtf_to_text
            text = rtf_to_text(text)
        parts = [text]
    text = "\n".join(parts).strip()
    if len(text) > MAX_TEXT:
        raise ValueError("Extracted text exceeds one million characters. Split the document first.")
    if not text:
        return Imported("", "No readable text was found. You can paste or type text manually, or import a clearer scan.")
    return Imported(text)


def lines_for_pages(text: str, width=90, height=45) -> list[list[str]]:
    lines = []
    for line in text.split("\n"):
        lines.extend(textwrap.wrap(line.expandtabs(4), width, replace_whitespace=False, drop_whitespace=False) or [""])
    return [lines[i:i+height] for i in range(0, len(lines), height)] or [[""]]


def export_file(path: Path, text: str) -> None:
    """Write only the supplied text; no source objects are accepted by this function."""
    extension = path.suffix.lower()
    if extension not in OUTPUT_EXTENSIONS:
        raise ValueError("Choose one of the supported export formats.")
    path.parent.mkdir(parents=True, exist_ok=True)
    if extension in {".txt", ".md", ".tsv"}:
        path.write_text(text, encoding="utf-8")
    elif extension == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, quoting=csv.QUOTE_ALL)
            for line in text.split("\n"):
                # Prevent spreadsheet formula execution when a CSV is opened.
                cells = line.split("\t")
                writer.writerow(["'" + c if c.lstrip().startswith(("=", "+", "-", "@")) else c for c in cells])
    elif extension == ".docx":
        from docx import Document
        doc = Document()
        doc.core_properties.author = "Redactor"
        doc.core_properties.title = "Reviewed work product"
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(path)
    elif extension == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        sheet = book.active
        sheet.title = "Work product"
        for row, line in enumerate(text.split("\n"), 1):
            for col, value in enumerate(line.split("\t"), 1):
                if len(value) > 32767 or col > 16384 or row > 1048576:
                    raise ValueError("Text exceeds Excel's cell or sheet limits. Export as TXT instead.")
                cell = sheet.cell(row, col, value)
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Calibri", size=11)
        sheet.column_dimensions["A"].width = 85
        book.properties.creator = "Redactor"
        book.save(path)
    elif extension == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        for page in lines_for_pages(text, width=85, height=19):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(.6), Inches(.5), Inches(12), Inches(6.5))
            frame = box.text_frame
            frame.word_wrap = True
            for index, line in enumerate(page):
                p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                p.text = line
                p.font.size = Pt(18)
        deck.core_properties.author = "Redactor"
        deck.save(path)
    elif extension == ".pdf":
        # Qt embeds a Unicode-capable system font and paginates the reviewed text.
        from PySide6.QtGui import QPdfWriter, QTextDocument, QFont, QPageSize, QPageLayout
        from PySide6.QtCore import QMarginsF, QBuffer, QIODevice, QSizeF
        destination = QBuffer()
        if not destination.open(QIODevice.OpenModeFlag.WriteOnly):
            raise OSError("R005: Unable to create PDF output buffer.")
        writer = QPdfWriter(destination)
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        writer.setPageMargins(QMarginsF(16, 16, 16, 16), QPageLayout.Unit.Millimeter)
        writer.setCreator("Redactor")
        writer.setTitle("Reviewed work product")
        doc = QTextDocument()
        doc.setDefaultFont(QFont("Segoe UI", 11))
        doc.setPlainText(text)
        # Explicit pagination prevents QTextDocument from injecting page numbers.
        doc.documentLayout().setPaintDevice(writer)
        doc.setPageSize(QSizeF(writer.width(), writer.height()))
        doc.print_(writer)
        del writer
        payload = bytes(destination.data())
        if not payload.startswith(b'%PDF-'):
            raise OSError("R005: PDF rendering failed; the destination was not changed.")
        import tempfile
        from .portable import atomic_replace
        handle, temporary = tempfile.mkstemp(dir=path.parent, prefix='.redactor-pdf-', suffix='.tmp')
        try:
            with os.fdopen(handle, 'wb') as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(stream.fileno())
            atomic_replace(temporary, path)
        finally:
            if os.path.exists(temporary): os.unlink(temporary)
    elif extension == ".rtf":
        def escape(value):
            out = []
            for c in value:
                if c in "\\{}": out.append("\\" + c)
                elif c == "\n": out.append("\\par\n")
                elif c == "\t": out.append("\\tab ")
                elif ord(c) < 128: out.append(c)
                else:
                    raw = c.encode("utf-16-le")
                    for i in range(0, len(raw), 2):
                        n = int.from_bytes(raw[i:i+2], "little", signed=True)
                        out.append(f"\\u{n}?")
            return "".join(out)
        path.write_text("{\\rtf1\\ansi\\uc1 " + escape(text) + "}", encoding="ascii")
    elif extension in {".odt", ".ods", ".odp"}:
        from odf.opendocument import OpenDocumentText, OpenDocumentSpreadsheet, OpenDocumentPresentation
        from odf.text import P
        if extension == ".odt":
            doc = OpenDocumentText()
            for line in text.split("\n"): doc.text.addElement(P(text=line))
        elif extension == ".ods":
            from odf.table import Table, TableRow, TableCell
            doc = OpenDocumentSpreadsheet()
            table = Table(name="Work product")
            for line in text.split("\n"):
                row = TableRow()
                for value in line.split("\t"):
                    cell = TableCell(valuetype="string")
                    cell.addElement(P(text=value))
                    row.addElement(cell)
                table.addElement(row)
            doc.spreadsheet.addElement(table)
        else:
            from odf.draw import Page, Frame, TextBox
            from odf.style import PageLayout, PageLayoutProperties, MasterPage
            doc = OpenDocumentPresentation()
            layout = PageLayout(name="Layout")
            layout.addElement(PageLayoutProperties(pagewidth="28cm", pageheight="21cm", printorientation="landscape"))
            doc.automaticstyles.addElement(layout)
            doc.masterstyles.addElement(MasterPage(name="Default", pagelayoutname="Layout"))
            for index, lines in enumerate(lines_for_pages(text, 80, 25)):
                page = Page(name=f"Page{index+1}", masterpagename="Default")
                frame = Frame(width="26cm", height="19cm", x="1cm", y="1cm")
                box = TextBox()
                for line in lines: box.addElement(P(text=line))
                frame.addElement(box)
                page.addElement(frame)
                doc.presentation.addElement(page)
        doc.save(str(path), addsuffix=False)
    else:
        from PIL import Image, ImageDraw, ImageFont
        pages = lines_for_pages(text, 85, 48)
        if len(pages) > 1 and extension not in {".tif", ".tiff"}:
            raise ValueError("Text needs multiple image pages. Export as TIFF or PDF to retain all content.")
        if len(pages) > 100:
            raise ValueError("Image export limit is 100 pages. Export as PDF or TXT instead.")
        font_paths = [Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "consola.ttf",
                      Path("/System/Library/Fonts/Menlo.ttc"), Path("/usr/share/fonts/dejavu-sans-mono-fonts/DejaVuSansMono.ttf"),
                      Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")]
        font_path = next((p for p in font_paths if p.exists()), None)
        font = ImageFont.truetype(str(font_path), 22) if font_path else ImageFont.load_default(size=22)
        images = []
        for lines in pages:
            image = Image.new("RGB", (1400, 1700), "white")
            draw = ImageDraw.Draw(image)
            for index, line in enumerate(lines):
                draw.text((55, 55 + index * 32), line, font=font, fill="#172b29")
            images.append(image)
        if extension in {".tif", ".tiff"}:
            images[0].save(path, save_all=True, append_images=images[1:], compression="tiff_deflate")
        else:
            images[0].save(path)
